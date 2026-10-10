# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.

"""PPTX: one block per text-bearing shape, plus one for the speaker notes.

Slides are numbered by their position in the presentation (1-based),
hidden slides included. Shapes are read in shape-tree order (the order
PowerPoint stores them, which is also their z-order); group shapes are
read recursively. Speaker notes are separate `slide_notes` blocks.

Not extracted in v0.1: charts, SmartArt, images (including alt text) and
slide masters/layouts.
"""

import io
import zipfile
from collections.abc import Iterable, Iterator

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide
from pptx.table import Table

from app.domain.enums import BlockKind
from app.ingestion.errors import IngestionError
from app.ingestion.guards import IngestionLimits
from app.ingestion.parsers.common import Parser, make_block, parser_version
from app.ingestion.schemas import ParsedBlock

_CELL_SEPARATOR = " | "
_OPEN_ERRORS = (PackageNotFoundError, KeyError, zipfile.BadZipFile, XMLSyntaxError, ValueError)


def parse_pptx(content: bytes, limits: IngestionLimits) -> list[ParsedBlock]:
    try:
        presentation = Presentation(io.BytesIO(content))
    except _OPEN_ERRORS:
        raise IngestionError.corrupt_file("not a readable PPTX presentation") from None

    blocks: list[ParsedBlock] = []
    for number, slide in enumerate(presentation.slides, start=1):
        for text in _shape_texts(slide.shapes):
            block = make_block(BlockKind.SLIDE_TEXT, text, slide_number=number)
            if block is not None:
                blocks.append(block)
        notes = make_block(BlockKind.SLIDE_NOTES, _notes_text(slide), slide_number=number)
        if notes is not None:
            blocks.append(notes)
    return blocks


def _shape_texts(shapes: Iterable[BaseShape]) -> Iterator[str]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text  # type: ignore[attr-defined]
        elif isinstance(shape, GraphicFrame) and shape.has_table:
            yield _table_text(shape.table)


def _table_text(table: Table) -> str:
    """Rows on separate lines, cells joined by " | ". Cells covered by a
    merge are skipped; the merged text is reported once, at its origin."""
    lines = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells if not cell.is_spanned]
        if any(cells):  # a row of empty cells is not text
            lines.append(_CELL_SEPARATOR.join(cells))
    return "\n".join(lines)


def _notes_text(slide: Slide) -> str:
    # Reading slide.notes_slide would create an empty notes slide; check first.
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text if frame is not None else ""


PPTX_PARSER = Parser(
    name="courseflow-pptx",
    version=parser_version(1, "python-pptx"),
    parse=parse_pptx,
    empty_reason="presentation contains no extractable text",
)
