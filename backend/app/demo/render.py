# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.

"""Render demo documents to PDF, DOCX, PPTX and HTML bytes, deterministically.

For fixed library versions the same document always gives the same bytes:
document metadata is pinned, ZIP members get a fixed timestamp and order,
and ReportLab runs in invariant mode (no creation date, no random document
ID). A library upgrade may still change the bytes; the byte-stability test
then fails and the files must be regenerated deliberately (a new dataset
version if any label depends on them).

ReportLab is a development dependency, imported only when a PDF is rendered:
the running application reads the committed files and never renders.
"""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime

from app.demo.northbridge_content import Document, DocxDoc, HtmlDoc, PdfDoc, PptxDoc

_FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_AUTHOR = "Northbridge University (fictional)"


def _pinned_zip(data: bytes) -> bytes:
    """Rewrite a ZIP with fixed member timestamps and attributes, same order."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            info = zipfile.ZipInfo(member.filename, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(member.filename))
    return out.getvalue()


def _pin_core_properties(properties: object, title: str) -> None:
    for name, value in {
        "author": _AUTHOR,
        "last_modified_by": _AUTHOR,
        "title": title,
        "created": _FIXED_TIME,
        "modified": _FIXED_TIME,
        "revision": 1,
    }.items():
        setattr(properties, name, value)


def render_pdf(doc: PdfDoc) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas

    fonts = {"title": ("Helvetica-Bold", 16, 22), "heading": ("Helvetica-Bold", 12, 18)}
    buffer = io.BytesIO()
    canvas = Canvas(buffer, pagesize=A4, invariant=1)
    canvas.setAuthor(_AUTHOR)
    canvas.setTitle(doc.title)
    for number, page in enumerate(doc.pages, start=1):
        y = 790.0
        for line in page:
            font, size, leading = fonts.get(line.style, ("Helvetica", 10, 14))
            if line.text:
                if canvas.stringWidth(line.text, font, size) > A4[0] - 2 * 56:
                    raise ValueError(f"{doc.title}: page {number} line is too wide: {line.text!r}")
                canvas.setFont(font, size)
                canvas.drawString(56, y, line.text)
            y -= leading
            if y < 60:
                raise ValueError(f"{doc.title}: page {number} overflows")
        canvas.setFont("Helvetica", 8)
        canvas.drawString(56, 30, f"Page {number} of {len(doc.pages)}")
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def render_docx(doc: DocxDoc) -> bytes:
    from docx import Document as WordDocument

    document = WordDocument()
    _pin_core_properties(document.core_properties, doc.title)
    styles = {"title": "Title", "heading1": "Heading 1", "heading2": "Heading 2"}
    for item in doc.items:
        if item.kind == "bullet":
            document.add_paragraph(item.text, style="List Bullet")
        elif item.kind in styles:
            document.add_paragraph(item.text, style=styles[item.kind])
        elif item.kind == "paragraph":
            document.add_paragraph(item.text)
        else:
            raise ValueError(f"unknown DOCX item kind {item.kind!r}")
    buffer = io.BytesIO()
    document.save(buffer)
    return _pinned_zip(buffer.getvalue())


_TITLE_SLIDE, _TITLE_AND_CONTENT = 0, 1


def render_pptx(doc: PptxDoc) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    _pin_core_properties(presentation.core_properties, doc.title)
    for index, spec in enumerate(doc.slides):
        layout = presentation.slide_layouts[_TITLE_SLIDE if index == 0 else _TITLE_AND_CONTENT]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = spec.title
        body = slide.placeholders[1].text_frame
        for position, bullet in enumerate(spec.bullets):
            paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
            paragraph.text = bullet
        if spec.notes is not None:
            slide.notes_slide.notes_text_frame.text = spec.notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return _pinned_zip(buffer.getvalue())


def render(doc: Document) -> bytes:
    if isinstance(doc, PdfDoc):
        return render_pdf(doc)
    if isinstance(doc, DocxDoc):
        return render_docx(doc)
    if isinstance(doc, PptxDoc):
        return render_pptx(doc)
    if isinstance(doc, HtmlDoc):
        return doc.html.encode("utf-8")
    raise TypeError(f"cannot render {type(doc).__name__}")
